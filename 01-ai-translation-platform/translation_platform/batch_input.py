"""批量翻译输入文件的只读加载。"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence, Tuple

from translation_platform.validation import InputValidationError


@dataclass(frozen=True)
class CellPosition:
    """源文件中一个单元格的一基位置。"""

    row_index: int
    column_index: int
    column_name: str


@dataclass(frozen=True)
class SourceCell:
    """待翻译的源文本及其原始位置。"""

    position: CellPosition
    text: str


@dataclass(frozen=True)
class BatchInput:
    """批量任务的源文件、表头和待翻译单元格。"""

    source_path: Path
    headers: Tuple[str, ...]
    selected_cells: Tuple[SourceCell, ...]


def load_batch_input(path: Path, columns: Sequence[str]) -> BatchInput:
    """只读加载 CSV 或 XLSX 中指定文本列的非空单元格。"""

    source_path = Path(path)
    normalized_columns = _normalize_columns(columns)
    if source_path.suffix.lower() == ".csv":
        headers, rows = _read_csv(source_path)
    elif source_path.suffix.lower() == ".xlsx":
        headers, rows = _read_xlsx(source_path)
    else:
        raise InputValidationError("输入文件必须为 .csv 或 .xlsx 格式")

    selected_indexes = _selected_indexes(headers, normalized_columns)
    selected_cells = _collect_selected_cells(rows, headers, selected_indexes)
    return BatchInput(source_path, headers, selected_cells)


def _normalize_columns(columns: Sequence[str]) -> Tuple[str, ...]:
    normalized = tuple(columns)
    if not normalized or any(not isinstance(column, str) or not column for column in normalized):
        raise InputValidationError("待翻译列必须包含至少一个非空文本列名")
    if len(set(normalized)) != len(normalized):
        raise InputValidationError("待翻译列不能包含重复列名")
    return normalized


def _read_csv(path: Path) -> Tuple[Tuple[str, ...], Tuple[Tuple[object, ...], ...]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream)
            header = next(reader, ())
            return _normalize_headers(header), tuple(tuple(row) for row in reader)
    except (OSError, UnicodeError, csv.Error) as exc:
        raise InputValidationError("无法读取 CSV 输入文件") from exc


def _read_xlsx(path: Path) -> Tuple[Tuple[str, ...], Tuple[Tuple[object, ...], ...]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise InputValidationError("读取 XLSX 输入文件需要安装 openpyxl 依赖") from exc

    try:
        with path.open("rb") as stream:
            workbook = load_workbook(stream, read_only=True, data_only=True)
            try:
                rows = workbook.active.iter_rows(values_only=True)
                header = next(rows, ())
                return _normalize_headers(header), tuple(tuple(row) for row in rows)
            finally:
                workbook.close()
    except Exception as exc:
        raise InputValidationError("无法读取 XLSX 输入文件") from exc


def _normalize_headers(values: Iterable[object]) -> Tuple[str, ...]:
    return tuple("" if value is None else str(value).strip() for value in values)


def _selected_indexes(headers: Tuple[str, ...], columns: Tuple[str, ...]) -> Tuple[int, ...]:
    missing = tuple(column for column in columns if column not in headers)
    if missing:
        raise InputValidationError("输入文件缺少列：" + "、".join(missing))
    return tuple(headers.index(column) for column in columns)


def _collect_selected_cells(
    rows: Iterable[Tuple[object, ...]],
    headers: Tuple[str, ...],
    selected_indexes: Tuple[int, ...],
) -> Tuple[SourceCell, ...]:
    selected_cells = []
    for row_index, row in enumerate(rows, start=2):
        for column_index in selected_indexes:
            value = row[column_index] if column_index < len(row) else None
            if value is None or value == "":
                continue
            if not isinstance(value, str):
                raise InputValidationError(
                    f"第 {row_index} 行第 {column_index + 1} 列的所选单元格必须为文本"
                )
            if not value.strip():
                continue
            selected_cells.append(
                SourceCell(
                    position=CellPosition(
                        row_index=row_index,
                        column_index=column_index + 1,
                        column_name=headers[column_index],
                    ),
                    text=value,
                )
            )
    return tuple(selected_cells)
