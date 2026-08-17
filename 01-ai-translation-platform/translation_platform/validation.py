"""命令行翻译请求的语义校验。"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path
from typing import Iterable, Tuple

from translation_platform.paths import resolve_project_path


SUPPORTED_FILE_EXTENSIONS = (".csv", ".xlsx")


class InputValidationError(ValueError):
    """用户提供的命令参数内部不一致时抛出。"""


def validate_arguments(arguments: argparse.Namespace) -> argparse.Namespace:
    """规范化并校验已解析的命令参数。"""

    _validate_language_pair(arguments.from_lang, arguments.to_lang)
    if arguments.command == "translate":
        arguments.text = _normalize_text(arguments.text)
    elif arguments.command == "batch":
        _validate_batch_arguments(arguments)
    return arguments


def _normalize_text(text: str) -> str:
    normalized = text.strip()
    if not normalized:
        raise InputValidationError("--text must not be empty")
    return normalized


def _validate_language_pair(from_lang: str, to_lang: str) -> None:
    if from_lang == to_lang:
        raise InputValidationError("source and target languages must differ")


def _validate_batch_arguments(arguments: argparse.Namespace) -> None:
    columns = _normalize_columns(arguments.columns)
    input_path = _normalize_path(arguments.input_path)
    output_path = _normalize_path(arguments.output_path)

    if input_path.suffix.lower() not in SUPPORTED_FILE_EXTENSIONS:
        raise InputValidationError("--input must be a .csv or .xlsx file")
    if not input_path.is_file():
        raise InputValidationError(f"input file does not exist: {input_path}")
    if output_path.suffix.lower() not in SUPPORTED_FILE_EXTENSIONS:
        raise InputValidationError("--output must be a .csv or .xlsx file")
    if input_path == output_path:
        raise InputValidationError("--output must differ from --input")

    available_columns = _read_header(input_path)
    missing_columns = tuple(column for column in columns if column not in available_columns)
    if missing_columns:
        raise InputValidationError(
            "input is missing columns: " + ", ".join(missing_columns)
        )

    arguments.input_path = input_path
    arguments.columns = columns
    arguments.output_path = output_path
    arguments.checkpoint_path = (
        _normalize_path(arguments.checkpoint_path)
        if arguments.checkpoint_path is not None
        else None
    )


def _normalize_columns(raw_columns: str) -> Tuple[str, ...]:
    columns = tuple(column.strip() for column in raw_columns.split(","))
    if not columns or any(not column for column in columns):
        raise InputValidationError("--columns contains an empty name")
    if len(set(columns)) != len(columns):
        raise InputValidationError("--columns contains duplicate names")
    return columns


def _normalize_path(raw_path: str) -> Path:
    return resolve_project_path(raw_path).resolve()


def _read_header(input_path: Path) -> Tuple[str, ...]:
    if input_path.suffix.lower() == ".csv":
        return _read_csv_header(input_path)
    return _read_xlsx_header(input_path)


def _read_csv_header(input_path: Path) -> Tuple[str, ...]:
    try:
        with input_path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream)
            header = next(reader, ())
    except (OSError, UnicodeError, csv.Error) as exc:
        raise InputValidationError(f"unable to read CSV header: {exc}") from exc
    return _normalize_header(header)


def _read_xlsx_header(input_path: Path) -> Tuple[str, ...]:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise InputValidationError(
            "XLSX validation requires the openpyxl dependency"
        ) from exc

    try:
        workbook = load_workbook(input_path, read_only=True, data_only=True)
        try:
            first_row = next(
                workbook.active.iter_rows(min_row=1, max_row=1, values_only=True),
                (),
            )
        finally:
            workbook.close()
    except (OSError, InvalidFileException) as exc:
        raise InputValidationError(f"unable to read XLSX header: {exc}") from exc
    return _normalize_header(first_row)


def _normalize_header(values: Iterable[object]) -> Tuple[str, ...]:
    return tuple("" if value is None else str(value).strip() for value in values)
