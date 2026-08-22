"""批量翻译结果的原子文件导出。"""

from __future__ import annotations

import csv
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Mapping, Sequence

from translation_platform.batch_input import CellPosition


class AtomicBatchExporter:
    """将指定单元格的译文安全写入新的 CSV 或 XLSX 文件。"""

    def write(
        self,
        source_path: Path,
        output_path: Path,
        replacements: Mapping[CellPosition, str],
    ) -> Path:
        """在输出目录创建临时文件，成功保存后原子替换目标文件。"""

        source = Path(source_path)
        output = Path(output_path)
        self._validate_paths(source, output)
        normalized_replacements = self._validate_replacements(replacements)
        temporary_path = self._create_temporary_path(output)
        try:
            if output.suffix.lower() == ".csv":
                self._write_csv(source, temporary_path, normalized_replacements)
            else:
                self._write_xlsx(source, temporary_path, normalized_replacements)
            self._sync_temporary_file(temporary_path)
            os.replace(temporary_path, output)
            return output
        except BaseException:
            self._remove_temporary_file(temporary_path)
            raise

    @staticmethod
    def _validate_paths(source: Path, output: Path) -> None:
        if source.resolve() == output.resolve():
            raise ValueError("输出路径不能与源文件相同")
        if source.suffix.lower() not in {".csv", ".xlsx"}:
            raise ValueError("源文件必须为 .csv 或 .xlsx 格式")
        if output.suffix.lower() not in {".csv", ".xlsx"}:
            raise ValueError("输出文件必须为 .csv 或 .xlsx 格式")

    @staticmethod
    def _validate_replacements(
        replacements: Mapping[CellPosition, str],
    ) -> Mapping[CellPosition, str]:
        normalized = dict(replacements)
        for position, text in normalized.items():
            if not isinstance(position, CellPosition):
                raise ValueError("译文位置必须是 CellPosition")
            if position.row_index < 1 or position.column_index < 1:
                raise ValueError("译文位置的行列坐标必须从 1 开始")
            if not isinstance(text, str):
                raise ValueError("译文必须为文本")
        return normalized

    @staticmethod
    def _create_temporary_path(output: Path) -> Path:
        with NamedTemporaryFile(
            prefix=".translation-platform-",
            suffix=output.suffix,
            dir=output.parent,
            delete=False,
        ) as temporary_file:
            return Path(temporary_file.name)

    @staticmethod
    def _write_csv(
        source: Path,
        temporary_path: Path,
        replacements: Mapping[CellPosition, str],
    ) -> None:
        if source.suffix.lower() == ".csv":
            rows = _read_csv_rows(source)
        else:
            rows = _read_xlsx_rows(source)
        replacements_by_row = _replacements_by_row(replacements)
        with temporary_path.open("w", encoding="utf-8-sig", newline="") as temporary_file:
            writer = csv.writer(temporary_file)
            for row_index, row in enumerate(rows, start=1):
                writer.writerow(_replace_csv_row(row, row_index, replacements_by_row))

    @staticmethod
    def _write_xlsx(
        source: Path,
        temporary_path: Path,
        replacements: Mapping[CellPosition, str],
    ) -> None:
        from openpyxl import Workbook, load_workbook

        if source.suffix.lower() == ".xlsx":
            workbook = load_workbook(source)
        else:
            workbook = Workbook()
            worksheet = workbook.active
            for row in _read_csv_rows(source):
                worksheet.append(row)
        try:
            worksheet = workbook.active
            for position, text in replacements.items():
                worksheet.cell(
                    row=position.row_index,
                    column=position.column_index,
                ).value = text
            workbook.save(temporary_path)
        finally:
            workbook.close()

    @staticmethod
    def _sync_temporary_file(temporary_path: Path) -> None:
        with temporary_path.open("r+b") as temporary_file:
            temporary_file.flush()
            os.fsync(temporary_file.fileno())

    @staticmethod
    def _remove_temporary_file(path: Path) -> None:
        try:
            path.unlink(missing_ok=True)
        except OSError:
            pass


def _replacements_by_row(
    replacements: Mapping[CellPosition, str],
) -> Mapping[int, Mapping[int, str]]:
    grouped: dict[int, dict[int, str]] = {}
    for position, text in replacements.items():
        grouped.setdefault(position.row_index, {})[position.column_index] = text
    return grouped


def _read_csv_rows(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as source_file:
        return list(csv.reader(source_file))


def _read_xlsx_rows(path: Path) -> list[tuple[object, ...]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()


def _replace_csv_row(
    row: Sequence[object],
    row_index: int,
    replacements_by_row: Mapping[int, Mapping[int, str]],
) -> list[str]:
    replaced_row = list(row)
    for column_index, text in replacements_by_row.get(row_index, {}).items():
        if column_index > len(replaced_row):
            raise ValueError(f"CSV 不包含第 {row_index} 行第 {column_index} 列")
        replaced_row[column_index - 1] = text
    return replaced_row
