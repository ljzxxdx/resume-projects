"""批量翻译结果的原子导出行为测试。"""

from __future__ import annotations

import csv
import hashlib
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from translation_platform.batch_input import CellPosition
from translation_platform.storage import AtomicBatchExporter


class AtomicBatchExporterTests(unittest.TestCase):
    """验证导出只修改新输出，并在失败时保留既有输出。"""

    def test_csv_writes_selected_cells_without_changing_source_or_other_values(self) -> None:
        """CSV 仅回填指定坐标，表头、未选中值和源文件均保持不变。"""

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "source.csv"
            output_path = Path(directory) / "translated.csv"
            _write_csv(
                source_path,
                [
                    ["编号", "标题", "说明"],
                    ["A-01", "原始标题", "原始说明"],
                    ["A-02", "保留标题", "保留说明"],
                ],
            )
            source_hash = _sha256(source_path)

            result = AtomicBatchExporter().write(
                source_path,
                output_path,
                {
                    CellPosition(2, 2, "标题"): "译文标题",
                    CellPosition(3, 3, "说明"): "译文说明",
                },
            )

            self.assertEqual(result, output_path)
            self.assertEqual(
                _read_csv(output_path),
                [
                    ["编号", "标题", "说明"],
                    ["A-01", "译文标题", "原始说明"],
                    ["A-02", "保留标题", "译文说明"],
                ],
            )
            self.assertEqual(source_hash, _sha256(source_path))
            self.assertEqual(
                _read_csv(source_path),
                [
                    ["编号", "标题", "说明"],
                    ["A-01", "原始标题", "原始说明"],
                    ["A-02", "保留标题", "保留说明"],
                ],
            )

    def test_xlsx_writes_selected_cells_without_changing_source_or_other_values(self) -> None:
        """XLSX 仅修改活动工作表的指定单元格，未选择值和源文件保持不变。"""

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "source.xlsx"
            output_path = Path(directory) / "translated.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "待翻译"
            worksheet.append(["编号", "标题", "说明"])
            worksheet.append(["A-01", "原始标题", "原始说明"])
            worksheet.append(["A-02", "保留标题", "保留说明"])
            workbook.create_sheet("不应修改").append(["原值"])
            workbook.save(source_path)
            source_hash = _sha256(source_path)

            result = AtomicBatchExporter().write(
                source_path,
                output_path,
                {
                    CellPosition(2, 2, "标题"): "译文标题",
                    CellPosition(3, 3, "说明"): "译文说明",
                },
            )

            self.assertEqual(result, output_path)
            exported_workbook = load_workbook(output_path)
            try:
                exported_sheet = exported_workbook.active
                self.assertEqual(exported_sheet["A1"].value, "编号")
                self.assertEqual(exported_sheet["B2"].value, "译文标题")
                self.assertEqual(exported_sheet["C2"].value, "原始说明")
                self.assertEqual(exported_sheet["B3"].value, "保留标题")
                self.assertEqual(exported_sheet["C3"].value, "译文说明")
                self.assertEqual(exported_workbook["不应修改"]["A1"].value, "原值")
            finally:
                exported_workbook.close()
            self.assertEqual(source_hash, _sha256(source_path))
            source_workbook = load_workbook(source_path)
            try:
                self.assertEqual(source_workbook.active["B2"].value, "原始标题")
                self.assertEqual(source_workbook.active["C3"].value, "保留说明")
            finally:
                source_workbook.close()

    def test_csv_input_writes_xlsx_output_without_changing_source(self) -> None:
        """CSV 输入可导出为 XLSX，且保留行列并仅替换指定单元格。"""

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "source.csv"
            output_path = Path(directory) / "translated.xlsx"
            _write_csv(
                source_path,
                [
                    ["编号", "标题", "说明"],
                    ["A-01", "原始标题", "原始说明"],
                    ["A-02", "保留标题", "保留说明"],
                ],
            )
            source_hash = _sha256(source_path)

            AtomicBatchExporter().write(
                source_path,
                output_path,
                {
                    CellPosition(2, 2, "标题"): "译文标题",
                    CellPosition(3, 3, "说明"): "译文说明",
                },
            )

            exported_workbook = load_workbook(output_path)
            try:
                exported_sheet = exported_workbook.active
                self.assertEqual(
                    list(exported_sheet.iter_rows(values_only=True)),
                    [
                        ("编号", "标题", "说明"),
                        ("A-01", "译文标题", "原始说明"),
                        ("A-02", "保留标题", "译文说明"),
                    ],
                )
            finally:
                exported_workbook.close()
            self.assertEqual(source_hash, _sha256(source_path))

    def test_xlsx_input_writes_csv_output_without_changing_source(self) -> None:
        """XLSX 输入可导出为 CSV，读取活动表并仅替换指定单元格。"""

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "source.xlsx"
            output_path = Path(directory) / "translated.csv"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["编号", "标题", "说明"])
            worksheet.append(["A-01", "原始标题", "原始说明"])
            worksheet.append(["A-02", "保留标题", "保留说明"])
            workbook.save(source_path)
            source_hash = _sha256(source_path)

            AtomicBatchExporter().write(
                source_path,
                output_path,
                {
                    CellPosition(2, 2, "标题"): "译文标题",
                    CellPosition(3, 3, "说明"): "译文说明",
                },
            )

            self.assertEqual(
                _read_csv(output_path),
                [
                    ["编号", "标题", "说明"],
                    ["A-01", "译文标题", "原始说明"],
                    ["A-02", "保留标题", "译文说明"],
                ],
            )
            self.assertEqual(source_hash, _sha256(source_path))

    def test_csv_and_xlsx_fsync_temporary_file_before_real_replace(self) -> None:
        """两种输出都必须真实同步临时文件，随后才执行真实原子替换。"""

        for suffix in (".csv", ".xlsx"):
            with self.subTest(suffix=suffix), tempfile.TemporaryDirectory() as directory:
                directory_path = Path(directory)
                source_path = directory_path / ("source" + suffix)
                output_path = directory_path / ("translated" + suffix)
                if suffix == ".csv":
                    _write_csv(source_path, [["标题"], ["原始标题"]])
                else:
                    _write_xlsx(source_path, "原始标题")

                events = []
                real_fsync = os.fsync
                real_replace = os.replace

                def recording_fsync(file_descriptor: int) -> None:
                    real_fsync(file_descriptor)
                    events.append("fsync")

                def recording_replace(source: Path, destination: Path) -> None:
                    self.assertTrue(Path(source).is_file())
                    real_replace(source, destination)
                    events.append("replace")

                with (
                    patch(
                        "translation_platform.storage.os.fsync",
                        side_effect=recording_fsync,
                    ),
                    patch(
                        "translation_platform.storage.os.replace",
                        side_effect=recording_replace,
                    ),
                ):
                    AtomicBatchExporter().write(
                        source_path,
                        output_path,
                        {CellPosition(2, 1, "标题"): "译文标题"},
                    )

                self.assertEqual(events, ["fsync", "replace"])
                if suffix == ".csv":
                    self.assertEqual(_read_csv(output_path), [["标题"], ["译文标题"]])
                else:
                    workbook = load_workbook(output_path, read_only=True)
                    try:
                        self.assertEqual(workbook.active["A2"].value, "译文标题")
                    finally:
                        workbook.close()

    def test_rejects_output_path_equal_to_source_before_writing(self) -> None:
        """输出路径与源文件相同会立即拒绝，防止源文件被覆盖。"""

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "source.csv"
            _write_csv(source_path, [["标题"], ["原始标题"]])
            source_hash = _sha256(source_path)

            with self.assertRaisesRegex(ValueError, "不能与源文件相同"):
                AtomicBatchExporter().write(
                    source_path,
                    source_path,
                    {CellPosition(2, 1, "标题"): "译文标题"},
                )

            self.assertEqual(source_hash, _sha256(source_path))
            self.assertEqual(_temporary_files(Path(directory)), [])

    def test_xlsx_save_failure_keeps_existing_output_and_removes_temporary_file(self) -> None:
        """工作簿实际保存失败时，旧输出字节不变且同目录临时文件被清理。"""

        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            source_path = directory_path / "source.xlsx"
            output_path = directory_path / "translated.xlsx"
            _write_xlsx(source_path, "原始标题")
            output_path.write_bytes(b"existing-output-bytes")
            original_output = output_path.read_bytes()
            source_hash = _sha256(source_path)

            with patch(
                "openpyxl.workbook.workbook.Workbook.save",
                side_effect=OSError("simulated save failure"),
            ):
                with self.assertRaisesRegex(OSError, "simulated save failure"):
                    AtomicBatchExporter().write(
                        source_path,
                        output_path,
                        {CellPosition(2, 1, "标题"): "译文标题"},
                    )

            self.assertEqual(output_path.read_bytes(), original_output)
            self.assertEqual(source_hash, _sha256(source_path))
            self.assertEqual(_temporary_files(directory_path), [])

    def test_replace_failure_keeps_existing_output_and_removes_temporary_file(self) -> None:
        """原子替换失败时，已真实写入的临时文件会清理且旧输出不受影响。"""

        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            source_path = directory_path / "source.csv"
            output_path = directory_path / "translated.csv"
            _write_csv(source_path, [["标题"], ["原始标题"]])
            output_path.write_bytes(b"existing-output-bytes")
            original_output = output_path.read_bytes()
            source_hash = _sha256(source_path)

            with patch(
                "translation_platform.storage.os.replace",
                side_effect=OSError("simulated replace failure"),
            ):
                with self.assertRaisesRegex(OSError, "simulated replace failure"):
                    AtomicBatchExporter().write(
                        source_path,
                        output_path,
                        {CellPosition(2, 1, "标题"): "译文标题"},
                    )

            self.assertEqual(output_path.read_bytes(), original_output)
            self.assertEqual(source_hash, _sha256(source_path))
            self.assertEqual(_temporary_files(directory_path), [])


def _write_csv(path: Path, rows: list[list[str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        csv.writer(stream).writerows(rows)


def _read_csv(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.reader(stream))


def _write_xlsx(path: Path, title: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["标题"])
    worksheet.append([title])
    workbook.save(path)


def _temporary_files(directory: Path) -> list[Path]:
    return list(directory.glob(".translation-platform-*"))


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
