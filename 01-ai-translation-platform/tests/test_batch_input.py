"""批量输入文件读取边界的行为测试。"""

from __future__ import annotations

import csv
import hashlib
import tempfile
import unittest
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from translation_platform.batch_input import (
    BatchInput,
    CellPosition,
    SourceCell,
    load_batch_input,
)
from translation_platform.validation import InputValidationError


class LoadBatchInputTests(unittest.TestCase):
    """验证输入读取会保留源单元格的位置且不改写源文件。"""

    def test_csv_returns_selected_text_cells_with_original_positions(self) -> None:
        """CSV 的多列文本保留表头名和一基行列坐标。"""

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "synthetic.csv"
            with source_path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerows(
                    [
                        ["编号", "标题", "说明"],
                        ["A-01", "合成标题", "合成说明"],
                        ["A-02", "", "第二条说明"],
                    ]
                )
            before_hash = _sha256(source_path)

            result = load_batch_input(source_path, ("标题", "说明"))

            self.assertEqual(
                result,
                BatchInput(
                    source_path=source_path,
                    headers=("编号", "标题", "说明"),
                    selected_cells=(
                        SourceCell(CellPosition(2, 2, "标题"), "合成标题"),
                        SourceCell(CellPosition(2, 3, "说明"), "合成说明"),
                        SourceCell(CellPosition(3, 3, "说明"), "第二条说明"),
                    ),
                ),
            )
            self.assertEqual(before_hash, _sha256(source_path))

    def test_xlsx_returns_selected_text_cells_with_original_positions(self) -> None:
        """XLSX 与 CSV 使用相同的单元格位置规则且只读源文件。"""

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "synthetic.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["编号", "标题", "说明"])
            worksheet.append(["A-01", "合成标题", "合成说明"])
            worksheet.append(["A-02", None, "第二条说明"])
            workbook.save(source_path)
            before_hash = _sha256(source_path)

            result = load_batch_input(source_path, ("标题", "说明"))

            self.assertEqual(
                result.headers,
                ("编号", "标题", "说明"),
            )
            self.assertEqual(
                result.selected_cells,
                (
                    SourceCell(CellPosition(2, 2, "标题"), "合成标题"),
                    SourceCell(CellPosition(2, 3, "说明"), "合成说明"),
                    SourceCell(CellPosition(3, 3, "说明"), "第二条说明"),
                ),
            )
            self.assertEqual(before_hash, _sha256(source_path))

    def test_csv_skips_selected_cells_containing_only_whitespace(self) -> None:
        """CSV 所选单元格规范化后为空时不创建请求候选。"""

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "whitespace.csv"
            with source_path.open("w", encoding="utf-8-sig", newline="") as stream:
                csv.writer(stream).writerows(
                    [["标题"], [" \t "], ["保留文本"]]
                )

            result = load_batch_input(source_path, ("标题",))

            self.assertEqual(
                result.selected_cells,
                (SourceCell(CellPosition(3, 1, "标题"), "保留文本"),),
            )

    def test_xlsx_skips_selected_cells_containing_only_whitespace(self) -> None:
        """XLSX 所选单元格规范化后为空时不创建请求候选。"""

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "whitespace.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["标题"])
            worksheet.append([" \t "])
            worksheet.append(["保留文本"])
            workbook.save(source_path)

            result = load_batch_input(source_path, ("标题",))

            self.assertEqual(
                result.selected_cells,
                (SourceCell(CellPosition(3, 1, "标题"), "保留文本"),),
            )

    def test_rejects_selected_xlsx_cell_that_is_not_text(self) -> None:
        """所选列含非文本单元格时拒绝创建翻译任务。"""

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "numeric.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["标题"])
            worksheet.append([7])
            workbook.save(source_path)

            with self.assertRaisesRegex(
                InputValidationError, "所选单元格必须为文本"
            ):
                load_batch_input(source_path, ("标题",))

    def test_rejects_unreadable_xlsx_input(self) -> None:
        """损坏的 XLSX 内容转换为输入校验异常。"""

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "broken.xlsx"
            source_path.write_bytes(b"not an xlsx file")

            with self.assertRaisesRegex(InputValidationError, "无法读取 XLSX 输入文件"):
                load_batch_input(source_path, ("标题",))

    def test_rejects_xlsx_with_malformed_internal_xml(self) -> None:
        """有效 ZIP 但内部 XML 损坏的 XLSX 也转换为输入校验异常。"""

        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "malformed.xlsx"
            workbook = Workbook()
            workbook.active.append(["标题"])
            workbook.save(source_path)
            rewritten_path = Path(directory) / "rewritten.xlsx"
            with ZipFile(source_path, "r") as source_archive:
                entries = [
                    (entry.filename, source_archive.read(entry))
                    for entry in source_archive.infolist()
                ]
            with ZipFile(rewritten_path, "w", compression=ZIP_DEFLATED) as archive:
                for filename, content in entries:
                    archive.writestr(
                        filename,
                        "<worksheet>" if filename == "xl/worksheets/sheet1.xml" else content,
                    )
            source_path.unlink()
            rewritten_path.rename(source_path)

            with self.assertRaisesRegex(InputValidationError, "无法读取 XLSX 输入文件"):
                load_batch_input(source_path, ("标题",))


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
