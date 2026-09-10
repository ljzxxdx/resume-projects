"""只读审计历史翻译日志。"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
from typing import Iterable

from openpyxl import load_workbook


_TIMESTAMP_PATTERN = re.compile(
    r"^\[(?P<timestamp>\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})\]"
)
_PAIR_PATTERN = re.compile(r"「(?P<input>.*?)」→「(?P<result>.*?)」$")
_WRITE_POSITION_PATTERN = re.compile(
    r"^\[\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}\]\s+"
    r"(?P<row>[1-9]\d*)行“(?P<column>[^”]+)”列\s*:"
)


@dataclass(frozen=True)
class RoundSummary:
    """一个由空行分隔的历史轮次摘要。"""

    round: int
    total_lines: int
    successful_lines: int
    failed_lines: int
    unparseable_lines: int

    def to_dict(self) -> dict[str, int]:
        """转换为可序列化摘要。"""
        return {
            "round": self.round,
            "total_lines": self.total_lines,
            "successful_lines": self.successful_lines,
            "failed_lines": self.failed_lines,
            "unparseable_lines": self.unparseable_lines,
        }


@dataclass(frozen=True)
class SuccessfulRecord:
    """可供后续核对复用的一条成功输入结果记录。"""

    line_number: int
    round: int
    input_text: str
    result_text: str
    row: int | None = None
    column: str | None = None


@dataclass(frozen=True)
class AuditSummary:
    """单份历史日志的统计结果。"""

    total_lines: int
    successful_lines: int
    failed_lines: int
    unique_inputs: int
    unique_input_result_pairs: int
    duplicates: int
    unparseable_lines: int
    rounds: tuple[RoundSummary, ...]

    def to_dict(self) -> dict[str, object]:
        """转换为可序列化摘要。"""
        return {
            "total_lines": self.total_lines,
            "successful_lines": self.successful_lines,
            "failed_lines": self.failed_lines,
            "unique_inputs": self.unique_inputs,
            "unique_input_result_pairs": self.unique_input_result_pairs,
            "duplicates": self.duplicates,
            "unparseable_lines": self.unparseable_lines,
            "rounds": [round_summary.to_dict() for round_summary in self.rounds],
        }


@dataclass(frozen=True)
class HistoryAudit:
    """两份历史日志的只读审计结果。"""

    translation_log: AuditSummary
    write_log: AuditSummary

    def to_dict(self) -> dict[str, object]:
        """转换为 CLI 使用的 JSON 结构。"""
        return {
            "translation_log": self.translation_log.to_dict(),
            "write_log": self.write_log.to_dict(),
        }


def _has_valid_timestamp(line: str) -> bool:
    """确认记录的时间戳格式正确。"""
    match = _TIMESTAMP_PATTERN.match(line)
    if match is None:
        return False
    try:
        datetime.strptime(match.group("timestamp"), "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return False
    return True


def _classify_line(line: str) -> tuple[str, tuple[str, str] | None]:
    """将非空行分类为成功、失败或无法解析。"""
    if not _has_valid_timestamp(line):
        return "unparseable", None
    match = _PAIR_PATTERN.search(line)
    if match is not None:
        return "success", (match.group("input"), match.group("result"))
    if "失败" in line:
        return "failure", None
    return "unparseable", None


def successful_records(lines: Iterable[str]) -> tuple[SuccessfulRecord, ...]:
    """解析成功输入结果对，供后续工作簿核对复用。"""
    records: list[SuccessfulRecord] = []
    round_number = 1
    has_lines_in_round = False
    for line_number, line in enumerate(lines, start=1):
        line = line.rstrip("\r\n")
        if not line:
            if has_lines_in_round:
                round_number += 1
                has_lines_in_round = False
            continue
        has_lines_in_round = True
        classification, pair = _classify_line(line)
        if classification == "success":
            assert pair is not None
            position = _WRITE_POSITION_PATTERN.match(line)
            records.append(
                SuccessfulRecord(
                    line_number=line_number,
                    round=round_number,
                    input_text=pair[0],
                    result_text=pair[1],
                    row=int(position.group("row")) if position is not None else None,
                    column=position.group("column") if position is not None else None,
                )
            )
    return tuple(records)


def audit_log_lines(lines: Iterable[str]) -> AuditSummary:
    """审计一份日志的行序列，不读取或修改外部文件。"""
    normalized_lines = tuple(line.rstrip("\r\n") for line in lines)
    successful_pairs: list[tuple[str, str]] = []
    successful_lines = 0
    failed_lines = 0
    unparseable_lines = 0
    rounds: list[RoundSummary] = []
    round_total = 0
    round_successful = 0
    round_failed = 0
    round_unparseable = 0

    def close_round() -> None:
        if round_total == 0:
            return
        rounds.append(
            RoundSummary(
                round=len(rounds) + 1,
                total_lines=round_total,
                successful_lines=round_successful,
                failed_lines=round_failed,
                unparseable_lines=round_unparseable,
            )
        )

    for line in normalized_lines:
        if not line:
            close_round()
            round_total = 0
            round_successful = 0
            round_failed = 0
            round_unparseable = 0
            continue

        round_total += 1
        classification, pair = _classify_line(line)
        if classification == "success":
            successful_lines += 1
            round_successful += 1
            assert pair is not None
            successful_pairs.append(pair)
        elif classification == "failure":
            failed_lines += 1
            round_failed += 1
        else:
            unparseable_lines += 1
            round_unparseable += 1
    close_round()

    unique_pairs = set(successful_pairs)
    return AuditSummary(
        total_lines=len(normalized_lines),
        successful_lines=successful_lines,
        failed_lines=failed_lines,
        unique_inputs=len({source for source, _ in successful_pairs}),
        unique_input_result_pairs=len(unique_pairs),
        duplicates=len(successful_pairs) - len(unique_pairs),
        unparseable_lines=unparseable_lines,
        rounds=tuple(rounds),
    )


def audit_file(path: Path) -> AuditSummary:
    """以只读方式审计一份 UTF-8 日志文件。"""
    if not path.is_file():
        raise FileNotFoundError("日志文件不存在或不是普通文件")
    with path.open("r", encoding="utf-8") as handle:
        return audit_log_lines(handle)


def audit_history(translation_log: Path, write_log: Path) -> HistoryAudit:
    """分别审计翻译日志和写入日志，不修改输入文件。"""
    return HistoryAudit(
        translation_log=audit_file(translation_log),
        write_log=audit_file(write_log),
    )


def _sha256_file(path: Path) -> str:
    """计算输入文件哈希，用于确认只读核对没有改动工作簿。"""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_successful_records(path: Path) -> tuple[SuccessfulRecord, ...]:
    """复用日志解析器读取成功记录，不向外暴露业务内容。"""
    with path.open("r", encoding="utf-8") as handle:
        return successful_records(handle)


def build_historical_evidence(
    *,
    translation_summary: dict[str, object],
    write_summary: dict[str, object],
    successful_write_records: int,
    verifiable_write_records: int,
    successful_matches: int,
    mismatches: int,
    unlocatable_records: int,
    unique_successful_write_cells: int,
) -> dict[str, object]:
    """构造不包含任何原文、译文或位置值的确定性聚合证据。"""
    unique_successful_texts = int(translation_summary.get("unique_inputs", 0))
    return {
        "schema_version": 1,
        "statistical_basis": {
            "log_unit": "日志物理行，包含空白轮次分隔行；日志行数不等于唯一翻译数。",
            "unique_text_unit": "翻译成功记录按输入文本去重。",
            "write_match_rule": "仅当日志中的逻辑行号和列名可定位，且目标单元格等于该记录结果时计为成功匹配。",
            "write_cell_unit": "成功匹配的工作簿单元格坐标去重后计数。",
            "unlocatable_rule": "缺少位置、工作表、列或逻辑行号无效时不猜测，计入无法定位。",
        },
        "logs": {
            "translation_log": translation_summary,
            "write_log": write_summary,
        },
        "workbook_reconciliation": {
            "successful_write_records": successful_write_records,
            "verifiable_write_records": verifiable_write_records,
            "successful_matches": successful_matches,
            "mismatches": mismatches,
            "unlocatable_records": unlocatable_records,
            "unique_successful_write_cells": unique_successful_write_cells,
        },
        "threshold": {
            "minimum": 1000,
            "rule": "成功唯一输入文本数或去重后的成功写入单元数至少达到 1000。",
            "is_met": (
                unique_successful_texts >= 1000
                or unique_successful_write_cells >= 1000
            ),
        },
    }


def reconcile_history(
    translation_log: Path,
    write_log: Path,
    workbook_path: Path,
    *,
    workbook_sheet: str | None = None,
) -> dict[str, object]:
    """只读核对写入日志与工作簿，并返回脱敏聚合证据。"""
    if not workbook_path.is_file():
        raise FileNotFoundError("工作簿不存在或不是普通文件")

    history = audit_history(translation_log, write_log)
    write_records = _read_successful_records(write_log)
    before_hash = _sha256_file(workbook_path)
    successful_matches = 0
    mismatches = 0
    unlocatable_records = 0
    successful_cells: set[tuple[int, int]] = set()

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if workbook_sheet is None:
            worksheet = workbook.worksheets[0] if workbook.worksheets else None
        else:
            worksheet = workbook[workbook_sheet] if workbook_sheet in workbook.sheetnames else None

        if worksheet is None:
            unlocatable_records = len(write_records)
        else:
            header_values = next(
                worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
                (),
            )
            header_columns: dict[str, list[int]] = {}
            for column_index, value in enumerate(header_values, start=1):
                if isinstance(value, str):
                    header_columns.setdefault(value, []).append(column_index)

            targets: list[tuple[SuccessfulRecord, int, int]] = []
            for record in write_records:
                if record.row is None or record.column is None:
                    unlocatable_records += 1
                    continue
                column_indexes = header_columns.get(record.column, [])
                physical_row = record.row + 1
                if (
                    len(column_indexes) != 1
                    or physical_row > worksheet.max_row
                    or physical_row < 2
                ):
                    unlocatable_records += 1
                    continue
                column_index = column_indexes[0]
                targets.append((record, physical_row, column_index))

            # read_only 工作表的随机取单元格会重复扫描文件，按目标坐标单次顺序读取。
            target_columns_by_row: dict[int, set[int]] = {}
            for _, row, column in targets:
                target_columns_by_row.setdefault(row, set()).add(column)
            target_values: dict[tuple[int, int], object] = {}
            if target_columns_by_row:
                for row_index, row_values in enumerate(
                    worksheet.iter_rows(min_row=2, values_only=True),
                    start=2,
                ):
                    for column_index in target_columns_by_row.get(row_index, ()):
                        target_values[(row_index, column_index)] = row_values[column_index - 1]

            for record, physical_row, column_index in targets:
                if target_values[(physical_row, column_index)] == record.result_text:
                    successful_matches += 1
                    successful_cells.add((physical_row, column_index))
                else:
                    mismatches += 1
    finally:
        workbook.close()

    after_hash = _sha256_file(workbook_path)
    if before_hash != after_hash:
        raise RuntimeError("只读核对后工作簿哈希发生变化")

    return build_historical_evidence(
        translation_summary=history.translation_log.to_dict(),
        write_summary=history.write_log.to_dict(),
        successful_write_records=len(write_records),
        verifiable_write_records=successful_matches + mismatches,
        successful_matches=successful_matches,
        mismatches=mismatches,
        unlocatable_records=unlocatable_records,
        unique_successful_write_cells=len(successful_cells),
    )


def write_historical_evidence(evidence: dict[str, object], output_path: Path) -> None:
    """原子替换聚合证据；替换失败时保留已有证据。"""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            "w",
            encoding="utf-8",
            dir=output_path.parent,
            prefix=f".{output_path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            temporary_path = Path(handle.name)
            json.dump(evidence, handle, ensure_ascii=False, indent=2, sort_keys=True)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_path, output_path)
    except OSError as error:
        raise OSError("原子写入聚合证据失败") from error
    finally:
        if temporary_path is not None and temporary_path.exists():
            try:
                temporary_path.unlink()
            except OSError:
                pass


def _build_parser() -> argparse.ArgumentParser:
    """创建命令行参数解析器。"""
    parser = argparse.ArgumentParser(description="只读审计历史翻译日志")
    parser.add_argument("--translation-log", type=Path, required=True)
    parser.add_argument("--write-log", type=Path, required=True)
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--workbook-sheet")
    parser.add_argument(
        "--evidence-output",
        type=Path,
        default=Path("artifacts/validation/historical_evidence.json"),
    )
    return parser


def main() -> int:
    """校验输入文件并将审计摘要输出为 JSON。"""
    parser = _build_parser()
    arguments = parser.parse_args()
    if not arguments.translation_log.is_file():
        parser.error("翻译日志文件不存在")
    if not arguments.write_log.is_file():
        parser.error("写入日志文件不存在")
    if arguments.workbook is None and arguments.workbook_sheet is not None:
        parser.error("指定工作表时必须同时提供工作簿")
    if arguments.workbook is not None and not arguments.workbook.is_file():
        parser.error("工作簿文件不存在")
    try:
        if arguments.workbook is None:
            report = audit_history(arguments.translation_log, arguments.write_log).to_dict()
        else:
            report = reconcile_history(
                arguments.translation_log,
                arguments.write_log,
                arguments.workbook,
                workbook_sheet=arguments.workbook_sheet,
            )
            write_historical_evidence(report, arguments.evidence_output)
    except UnicodeDecodeError:
        parser.error("日志文件不是有效的 UTF-8 文本")
    except FileNotFoundError:
        parser.error("日志文件状态发生变化，请重新执行")
    except OSError:
        parser.error("无法读取工作簿或原子写入聚合证据")
    print(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
