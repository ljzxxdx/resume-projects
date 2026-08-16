"""Tests for collection record persistence."""

from __future__ import annotations

import csv
import json
import tempfile
import threading
import time
import unittest
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from taobao_collector import exporters
from taobao_collector.collectors.base import CollectionResult
from taobao_collector.models import (
    CommentRecord,
    Engine,
    LiveRoomRecord,
    ProductRecord,
    RecordStatus,
    RunMode,
)


class CollectionExporterTests(unittest.TestCase):
    @staticmethod
    def make_single_product_result(run_id: str) -> CollectionResult:
        return CollectionResult(
            products=(
                ProductRecord(
                    run_id=run_id,
                    engine=Engine.DRISSION,
                    keyword="陶瓷",
                    product_id="atomic-1",
                    source_url=(
                        "https://item.taobao.com/item.htm?id=atomic-1"
                    ),
                    name="原子写入商品",
                ),
            )
        )

    def test_export_deduplicates_within_and_across_runs_with_stats(
        self,
    ) -> None:
        def make_result(run_id: str) -> CollectionResult:
            product = ProductRecord(
                run_id=run_id,
                engine=Engine.DRISSION,
                keyword="陶瓷",
                product_id="1001",
                source_url="https://item.taobao.com/item.htm?id=1001",
                name="茶杯",
            )
            comment = CommentRecord(
                run_id=run_id,
                engine=Engine.DRISSION,
                keyword="陶瓷",
                product_id="1001",
                source_url="https://item.taobao.com/item.htm?id=1001",
                user_name="用户甲",
                sku_info="白色",
                content="做工细致",
                status=RecordStatus.FAILED,
                error_message="字段不完整",
            )
            missing_identity = ProductRecord(
                run_id=run_id,
                engine=Engine.DRISSION,
                keyword="陶瓷",
                product_id="",
                source_url="",
                name="未知商品",
            )
            return CollectionResult(
                products=(product, product, missing_identity),
                comments=(comment, comment),
            )

        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)
            exporters.write_collection_result(
                make_result("run-first"),
                output_root,
                run_id="run-first",
                mode=RunMode.PRODUCTS,
            )
            first_stats_path = (
                output_root / "run-first" / "export_stats.json"
            )
            quality_path = (
                output_root / "run-first" / "data_quality.json"
            )
            self.assertTrue(
                first_stats_path.exists(),
                "导出尚未生成去重统计",
            )
            self.assertTrue(
                quality_path.exists(),
                "导出尚未生成数据质量报告",
            )
            first_stats = json.loads(
                first_stats_path.read_text(encoding="utf-8")
            )
            first_products = (
                output_root / "run-first" / "products.jsonl"
            ).read_text(encoding="utf-8").splitlines()
            first_comments = (
                output_root / "run-first" / "comments.jsonl"
            ).read_text(encoding="utf-8").splitlines()
            first_quality = json.loads(
                quality_path.read_text(encoding="utf-8")
            )

            exporters.write_collection_result(
                make_result("run-second"),
                output_root,
                run_id="run-second",
                mode=RunMode.PRODUCTS,
            )
            second_stats = json.loads(
                (output_root / "run-second" / "export_stats.json")
                .read_text(encoding="utf-8")
            )
            second_quality = json.loads(
                (output_root / "run-second" / "data_quality.json")
                .read_text(encoding="utf-8")
            )
            second_products = (
                output_root / "run-second" / "products.jsonl"
            ).read_text(encoding="utf-8")
            second_comments = (
                output_root / "run-second" / "comments.jsonl"
            ).read_text(encoding="utf-8")
            key_index = (output_root / ".dedup_keys.json").read_text(
                encoding="utf-8"
            )

        self.assertEqual(len(first_products), 1)
        self.assertEqual(len(first_comments), 1)
        self.assertEqual(
            first_quality["tables"]["products"]["record_count"],
            3,
        )
        self.assertEqual(
            first_quality["tables"]["products"]["unique_count"],
            1,
        )
        self.assertEqual(
            first_quality["tables"]["comments"]
            ["abnormal_status_count"],
            2,
        )
        self.assertEqual(
            first_stats,
            {
                "new_count": 2,
                "success_count": 1,
                "duplicate_count": 2,
                "missing_count": 1,
                "failed_count": 1,
                "tables": {
                    "products": {
                        "new_count": 1,
                        "success_count": 1,
                        "duplicate_count": 1,
                        "missing_count": 1,
                        "failed_count": 0,
                    },
                    "comments": {
                        "new_count": 1,
                        "success_count": 0,
                        "duplicate_count": 1,
                        "missing_count": 0,
                        "failed_count": 1,
                    },
                },
            },
        )
        self.assertEqual(second_products, "")
        self.assertEqual(second_comments, "")
        self.assertEqual(second_stats["new_count"], 0)
        self.assertEqual(second_stats["success_count"], 0)
        self.assertEqual(second_stats["duplicate_count"], 4)
        self.assertEqual(second_stats["missing_count"], 1)
        self.assertEqual(second_stats["failed_count"], 1)
        self.assertEqual(
            second_quality["tables"]["products"]["unique_count"],
            0,
        )
        self.assertEqual(
            second_quality["tables"]["products"]["duplicate_count"],
            2,
        )
        self.assertNotIn("用户甲", key_index)
        self.assertNotIn("做工细致", key_index)

    def test_products_mode_writes_independent_linked_tables_in_all_formats(
        self,
    ) -> None:
        collected_at = datetime(
            2026,
            7,
            30,
            10,
            0,
            tzinfo=timezone.utc,
        )
        result = CollectionResult(
            products=(
                ProductRecord(
                    run_id="run-export",
                    engine=Engine.DRISSION,
                    keyword="陶瓷",
                    product_id="1001",
                    source_url="https://item.taobao.com/item.htm?id=1001",
                    name="茶杯",
                    price=Decimal("39.90"),
                    sales_count=88,
                    collected_at=collected_at,
                ),
            ),
            comments=(
                CommentRecord(
                    run_id="run-export",
                    engine=Engine.DRISSION,
                    keyword="陶瓷",
                    product_id="1001",
                    source_url="https://item.taobao.com/item.htm?id=1001",
                    user_name="用户甲",
                    sku_info="白色",
                    content="做工细致",
                    collected_at=collected_at,
                ),
            ),
        )

        with tempfile.TemporaryDirectory() as directory:
            paths = exporters.write_collection_result(
                result,
                Path(directory),
                run_id="run-export",
                mode=RunMode.PRODUCTS,
            )

            self.assertEqual(
                {path.name for path in paths},
                {
                    "products.jsonl",
                    "products.csv",
                    "products.xlsx",
                    "comments.jsonl",
                    "comments.csv",
                    "comments.xlsx",
                },
            )
            product = json.loads(
                (Path(directory) / "run-export" / "products.jsonl")
                .read_text(encoding="utf-8")
                .strip()
            )
            comment = json.loads(
                (Path(directory) / "run-export" / "comments.jsonl")
                .read_text(encoding="utf-8")
                .strip()
            )
            with (
                Path(directory) / "run-export" / "comments.csv"
            ).open(encoding="utf-8-sig", newline="") as stream:
                csv_comment = next(csv.DictReader(stream))
            workbook = load_workbook(
                Path(directory) / "run-export" / "comments.xlsx",
                read_only=True,
            )
            sheet = workbook["comments"]
            rows = list(sheet.iter_rows(values_only=True))
            excel_comment = dict(zip(rows[0], rows[1]))
            workbook.close()

        self.assertEqual(product["price"], "39.90")
        self.assertEqual(product["engine"], "drission")
        self.assertEqual(
            product["collected_at"],
            "2026-07-30T10:00:00+00:00",
        )
        self.assertEqual(comment["content"], "做工细致")
        self.assertEqual(comment["product_id"], product["product_id"])
        self.assertEqual(comment["run_id"], product["run_id"])
        self.assertEqual(csv_comment["unique_key"], comment["unique_key"])
        self.assertEqual(excel_comment["unique_key"], comment["unique_key"])
        for repeated_product_field in ("name", "price", "sales_count"):
            self.assertNotIn(repeated_product_field, comment)
            self.assertNotIn(repeated_product_field, csv_comment)
            self.assertNotIn(repeated_product_field, excel_comment)

    def test_index_cache_failure_keeps_complete_run_and_manifest_dedup(
        self,
    ) -> None:
        original_replace = Path.replace

        def fail_index_replace(path: Path, target: Path):
            if Path(target).name == ".dedup_keys.json":
                raise OSError("模拟索引发布失败")
            return original_replace(path, target)

        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)
            with patch.object(Path, "replace", fail_index_replace):
                exporters.write_collection_result(
                    self.make_single_product_result("run-atomic"),
                    output_root,
                    run_id="run-atomic",
                    mode=RunMode.PRODUCTS,
                )

            first_run = output_root / "run-atomic"
            self.assertTrue((first_run / "products.jsonl").exists())
            self.assertTrue((first_run / ".unique_keys.json").exists())
            self.assertFalse((output_root / ".dedup_keys.json").exists())

            exporters.write_collection_result(
                self.make_single_product_result("run-after-cache-failure"),
                output_root,
                run_id="run-after-cache-failure",
                mode=RunMode.PRODUCTS,
            )
            self.assertEqual(
                (
                    output_root
                    / "run-after-cache-failure"
                    / "products.jsonl"
                ).read_text(encoding="utf-8"),
                "",
            )

    def test_malformed_key_index_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)
            (output_root / ".dedup_keys.json").write_text(
                json.dumps({"products": "product:id:atomic-1"}),
                encoding="utf-8",
            )

            with self.assertRaisesRegex(ValueError, "去重索引格式无效"):
                exporters.write_collection_result(
                    self.make_single_product_result("run-corrupt"),
                    output_root,
                    run_id="run-corrupt",
                    mode=RunMode.PRODUCTS,
                )

            self.assertFalse((output_root / "run-corrupt").exists())

    def test_key_index_with_missing_table_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)
            (output_root / ".dedup_keys.json").write_text(
                json.dumps({"products": []}),
                encoding="utf-8",
            )

            with self.assertRaisesRegex(ValueError, "去重索引格式无效"):
                exporters.write_collection_result(
                    self.make_single_product_result("run-missing-table"),
                    output_root,
                    run_id="run-missing-table",
                    mode=RunMode.PRODUCTS,
                )

            self.assertFalse(
                (output_root / "run-missing-table").exists()
            )

    def test_preexisting_unlocked_lock_file_does_not_block_export(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)
            (output_root / ".export.lock").write_text(
                "stale lock metadata",
                encoding="utf-8",
            )

            exporters.write_collection_result(
                self.make_single_product_result("run-after-crash"),
                output_root,
                run_id="run-after-crash",
                mode=RunMode.PRODUCTS,
            )

            self.assertTrue(
                (output_root / "run-after-crash" / "products.jsonl")
                .exists()
            )

    def test_unpublished_temporary_manifest_is_not_dedup_authority(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)
            abandoned = output_root / ".run-crashed-123"
            abandoned.mkdir()
            abandoned.joinpath(".unique_keys.json").write_text(
                json.dumps(
                    {
                        "products": ["product:id:atomic-1"],
                        "comments": [],
                        "live_rooms": [],
                    }
                ),
                encoding="utf-8",
            )

            exporters.write_collection_result(
                self.make_single_product_result("run-after-temp-crash"),
                output_root,
                run_id="run-after-temp-crash",
                mode=RunMode.PRODUCTS,
            )

            lines = (
                output_root
                / "run-after-temp-crash"
                / "products.jsonl"
            ).read_text(encoding="utf-8").splitlines()

        self.assertEqual(len(lines), 1)

    def test_concurrent_runs_publish_only_one_copy_of_same_record(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)
            first_loaded = threading.Event()
            release_first = threading.Event()
            errors = []
            original_load = exporters._load_key_index

            def delayed_load(path: Path):
                loaded = original_load(path)
                if threading.current_thread().name == "first-export":
                    first_loaded.set()
                    release_first.wait(timeout=2)
                return loaded

            def export(run_id: str) -> None:
                try:
                    exporters.write_collection_result(
                        self.make_single_product_result(run_id),
                        output_root,
                        run_id=run_id,
                        mode=RunMode.PRODUCTS,
                    )
                except Exception as exc:
                    errors.append(exc)

            with patch.object(
                exporters,
                "_load_key_index",
                delayed_load,
            ):
                first = threading.Thread(
                    target=export,
                    args=("run-concurrent-a",),
                    name="first-export",
                )
                second = threading.Thread(
                    target=export,
                    args=("run-concurrent-b",),
                    name="second-export",
                )
                first.start()
                self.assertTrue(first_loaded.wait(timeout=2))
                second.start()
                time.sleep(0.1)
                release_first.set()
                first.join(timeout=3)
                second.join(timeout=3)

            total_lines = sum(
                len(
                    (output_root / run_id / "products.jsonl")
                    .read_text(encoding="utf-8")
                    .splitlines()
                )
                for run_id in ("run-concurrent-a", "run-concurrent-b")
            )

        self.assertEqual(errors, [])
        self.assertFalse(first.is_alive())
        self.assertFalse(second.is_alive())
        self.assertEqual(total_lines, 1)

    def test_live_mode_creates_empty_file_when_no_rooms_are_found(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            paths = exporters.write_collection_result(
                CollectionResult(),
                Path(directory),
                run_id="run-empty-live",
                mode=RunMode.LIVE,
            )

            live_path = Path(directory) / "run-empty-live" / "live_rooms.jsonl"
            self.assertEqual(
                {path.name for path in paths},
                {
                    "live_rooms.jsonl",
                    "live_rooms.csv",
                    "live_rooms.xlsx",
                },
            )
            self.assertTrue(live_path.exists())
            self.assertEqual(live_path.read_text(encoding="utf-8"), "")
            with (
                Path(directory) / "run-empty-live" / "live_rooms.csv"
            ).open(encoding="utf-8-sig", newline="") as stream:
                self.assertIn("live_room_id", next(csv.reader(stream)))
            workbook = load_workbook(
                Path(directory) / "run-empty-live" / "live_rooms.xlsx",
                read_only=True,
            )
            self.assertEqual(workbook.active.title, "live_rooms")
            self.assertIn(
                "live_room_id",
                next(workbook.active.iter_rows(values_only=True)),
            )
            workbook.close()

    def test_quality_ignores_records_outside_the_selected_mode(self) -> None:
        result = self.make_single_product_result("run-live-only")

        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)
            exporters.write_collection_result(
                result,
                output_root,
                run_id="run-live-only",
                mode=RunMode.LIVE,
            )
            quality = json.loads(
                (output_root / "run-live-only" / "data_quality.json")
                .read_text(encoding="utf-8")
            )

        self.assertEqual(
            quality["tables"]["products"],
            {
                "record_count": 0,
                "unique_count": 0,
                "duplicate_count": 0,
                "required_field_missing_count": 0,
                "required_field_missing_by_field": {},
                "abnormal_status_count": 0,
            },
        )

    def test_live_quality_counts_missing_identity_and_bad_status(
        self,
    ) -> None:
        result = CollectionResult(
            live_rooms=(
                LiveRoomRecord(
                    run_id="run-live-quality",
                    engine=Engine.DRISSION,
                    keyword="陶瓷",
                    live_room_id="",
                    source_url="",
                    account_name="",
                    introduction="页面字段缺失",
                    status=RecordStatus.BLOCKED,
                ),
            )
        )

        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)
            exporters.write_collection_result(
                result,
                output_root,
                run_id="run-live-quality",
                mode=RunMode.LIVE,
            )
            quality = json.loads(
                (output_root / "run-live-quality" / "data_quality.json")
                .read_text(encoding="utf-8")
            )["tables"]["live_rooms"]

        self.assertEqual(quality["record_count"], 1)
        self.assertEqual(quality["unique_count"], 0)
        self.assertEqual(quality["required_field_missing_count"], 3)
        self.assertEqual(quality["abnormal_status_count"], 1)

    def test_spreadsheet_outputs_escape_formula_like_collected_text(
        self,
    ) -> None:
        result = CollectionResult(
            products=(
                ProductRecord(
                    run_id="run-formula",
                    engine=Engine.DRISSION,
                    keyword="陶瓷",
                    product_id="formula-1",
                    source_url="https://item.taobao.com/item.htm?id=formula-1",
                    name="=HYPERLINK(\"https://example.test\")",
                ),
            )
        )

        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)
            exporters.write_collection_result(
                result,
                output_root,
                run_id="run-formula",
                mode=RunMode.PRODUCTS,
            )
            with (
                output_root / "run-formula" / "products.csv"
            ).open(encoding="utf-8-sig", newline="") as stream:
                csv_product = next(csv.DictReader(stream))
            workbook = load_workbook(
                output_root / "run-formula" / "products.xlsx",
                read_only=True,
                data_only=False,
            )
            rows = list(workbook["products"].iter_rows(values_only=True))
            excel_product = dict(zip(rows[0], rows[1]))
            workbook.close()

        self.assertTrue(csv_product["name"].startswith("'="))
        self.assertTrue(excel_product["name"].startswith("'="))

    def test_export_writes_redacted_sample_with_approved_comment_only(
        self,
    ) -> None:
        approved = CommentRecord(
            run_id="run-redacted",
            engine=Engine.DRISSION,
            keyword="陶瓷",
            product_id="1001",
            source_url="https://item.taobao.com/item.htm?id=1001",
            user_name="用户甲乙",
            sku_info="白色",
            content="允许公开的评论",
        )
        private = CommentRecord(
            run_id="run-redacted",
            engine=Engine.DRISSION,
            keyword="陶瓷",
            product_id="1001",
            source_url="https://item.taobao.com/item.htm?id=1001",
            user_name="用户丙丁",
            sku_info="蓝色",
            content="不允许公开的评论",
        )

        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)
            exporters.write_collection_result(
                CollectionResult(
                    comments=(approved, approved, private),
                ),
                output_root,
                run_id="run-redacted",
                mode=RunMode.PRODUCTS,
                approved_comment_keys={
                    exporters.record_unique_key(approved)
                },
            )
            sample_text = (
                output_root / "run-redacted" / "redacted_sample.json"
            ).read_text(encoding="utf-8")
            sample = json.loads(sample_text)

        self.assertEqual(len(sample["comments"]), 1)
        self.assertEqual(
            sample["comments"][0]["content"],
            "允许公开的评论",
        )
        self.assertNotIn("不允许公开的评论", sample_text)
        self.assertNotIn("1001", sample_text)

    def test_historical_comment_can_be_published_after_manual_approval(
        self,
    ) -> None:
        comment = CommentRecord(
            run_id="run-first-private",
            engine=Engine.DRISSION,
            keyword="陶瓷",
            product_id="1001",
            source_url="https://item.taobao.com/item.htm?id=1001",
            user_name="用户甲乙",
            sku_info="白色",
            content="首轮后人工确认公开",
        )
        approved_key = exporters.record_unique_key(comment)

        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)
            exporters.write_collection_result(
                CollectionResult(comments=(comment,)),
                output_root,
                run_id="run-first-private",
                mode=RunMode.PRODUCTS,
            )
            first_sample = json.loads(
                (
                    output_root
                    / "run-first-private"
                    / "redacted_sample.json"
                ).read_text(encoding="utf-8")
            )

            exporters.write_collection_result(
                CollectionResult(comments=(comment, comment)),
                output_root,
                run_id="run-second-approved",
                mode=RunMode.PRODUCTS,
                approved_comment_keys={approved_key},
            )
            second_sample = json.loads(
                (
                    output_root
                    / "run-second-approved"
                    / "redacted_sample.json"
                ).read_text(encoding="utf-8")
            )
            second_comments = (
                output_root
                / "run-second-approved"
                / "comments.jsonl"
            ).read_text(encoding="utf-8")

        self.assertEqual(first_sample["comments"], [])
        self.assertEqual(second_comments, "")
        self.assertEqual(len(second_sample["comments"]), 1)
        self.assertEqual(
            second_sample["comments"][0]["content"],
            "首轮后人工确认公开",
        )

    def test_products_mode_does_not_publish_partial_files(self) -> None:
        valid_product = ProductRecord(
            run_id="run-partial",
            engine=Engine.DRISSION,
            keyword="陶瓷",
            product_id="1001",
            source_url="https://item.taobao.com/1001",
            name="茶杯",
        )
        invalid_result = CollectionResult(
            products=(valid_product,),
            comments=(object(),),
        )

        with tempfile.TemporaryDirectory() as directory:
            output_root = Path(directory)

            with self.assertRaises(TypeError):
                exporters.write_collection_result(
                    invalid_result,
                    output_root,
                    run_id="run-partial",
                    mode=RunMode.PRODUCTS,
                )

            self.assertFalse((output_root / "run-partial").exists())
            self.assertEqual(
                {path.name for path in output_root.iterdir()},
                {".export.lock"},
            )


if __name__ == "__main__":
    unittest.main()
